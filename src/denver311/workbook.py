"""Build the analyst workbook (reports/denver311_workbook.xlsx) from the aggregate exports.

The same file is imported into Google Sheets and opens unchanged in Excel. Aggregates only:
no addresses, no row-level records.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from denver311 import EXPORT_DIR, ROOT

OUT = ROOT / "reports" / "denver311_workbook.xlsx"
REPO_URL = "https://github.com/jasonpellerin-aisolutionist/denver-311-service-analytics"
SOURCE_URL = (
    "https://www.denvergov.org/opendata/dataset/city-and-county-of-denver-311-service-requests"
)

NAVY = "1E3A5F"
SLATE = "64748B"
TINT = "EEF2F7"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="0F172A")
NOTE_FONT = Font(italic=True, color=SLATE)


def export(name: str) -> pd.DataFrame:
    return pd.read_csv(EXPORT_DIR / f"{name}.csv")


def write_table(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    formats: dict[str, str] | None = None,
    widths: dict[str, int] | None = None,
) -> None:
    formats = formats or {}
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            value = None if pd.isna(value) else value
            cell = ws.cell(row=i, column=j, value=value.item() if hasattr(value, "item") else value)
            fmt = formats.get(df.columns[j - 1])
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for j, col in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        width = (widths or {}).get(col) or min(
            max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2, 60
        )
        ws.column_dimensions[letter].width = width


def color_scale(ws, col_letter: str, first: int, last: int) -> None:
    ws.conditional_formatting.add(
        f"{col_letter}{first}:{col_letter}{last}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="F8FAFC",
            end_type="num",
            end_value=1,
            end_color="1E3A5F",
        ),
    )


def readme(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "Denver 311 Service Analytics, 2019-2025"
    ws["A1"].font = TITLE_FONT
    lines = [
        "Question: where does Denver 311 demand come from, how fast does it close, and what would an operations lead change?",
        "Source: Denver Open Data Catalog, 311 Service Requests, 2019-2025 yearly files (CC BY 3.0).",
        "Neighborhoods: ACS 2019-2023 statistical neighborhood polygons, Denver Open Data Catalog.",
        "Contents: aggregates only. No addresses or row-level records.",
        "Benchmark: each request type's own 2019 P90 time to close. Derived for this analysis, not a city service-level target.",
        "Published goals: DOTI Level of Service goals (business days) for the five top types that have one. See the Published LOS tab.",
        "Wait index: rank of each closed request against the same type and year (0 fastest, 1 slowest). 0.50 is typical.",
        "Record gap: over 20% of a type-year never closed in the file, or over 25% of its closures stamped on one day.",
        "Built with Python (pandas), DuckDB, PostgreSQL. Analysis: Jason Pellerin.",
    ]
    for i, text in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=text)
    ws.cell(row=len(lines) + 4, column=1, value="Code and method").font = Font(bold=True)
    link = ws.cell(row=len(lines) + 5, column=1, value=REPO_URL)
    link.hyperlink, link.font = REPO_URL, Font(color="2563EB", underline="single")
    src = ws.cell(row=len(lines) + 6, column=1, value=SOURCE_URL)
    src.hyperlink, src.font = SOURCE_URL, Font(color="2563EB", underline="single")
    ws.column_dimensions["A"].width = 120


def kpis(wb: Workbook) -> None:
    ws = wb.create_sheet("KPI Summary")
    df = export("kpi_summary")
    write_table(ws, df, widths={"metric": 44, "note": 90})
    for i, unit in enumerate(df["unit"], start=2):
        ws.cell(row=i, column=2).number_format = {
            "count": "#,##0",
            "share": "0.0%",
            "change": "+0.0%;-0.0%",
            "index": "0.000",
        }.get(unit, "General")


def dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Dictionary")
    write_table(ws, pd.read_csv(ROOT / "data" / "data_dictionary.csv"), widths={"description": 90})


def quality(wb: Workbook) -> None:
    ws = wb.create_sheet("Quality Log")
    ws["A1"] = "Cleaning steps (src/denver311/clean.py)"
    ws["A1"].font = Font(bold=True, size=12)
    log = export("cleaning_log")
    write_table(
        ws,
        log,
        start_row=2,
        formats={"rows": "#,##0"},
        widths={"step": 30, "rows": 12, "decision": 100},
    )
    start = len(log) + 5
    ws.cell(
        row=start - 1, column=1, value="Post-load checks (sql/duckdb/02_quality_checks.sql)"
    ).font = Font(bold=True, size=12)
    write_table(
        ws,
        export("quality_checks"),
        start_row=start,
        formats={"failing_rows": "#,##0"},
        widths={"check_name": 30, "failing_rows": 12, "tolerance": 100},
    )
    ws.freeze_panes = "A3"


def by_type(wb: Workbook) -> None:
    ws = wb.create_sheet("KPI by Type")
    df = export("type_year")
    df = df.sort_values(["requests_total", "request_type", "year"], ascending=[False, True, True])
    df = pd.DataFrame(
        {
            "request_type": df.request_type,
            "category": df.request_category,
            "agency": df.agency,
            "year": df.year,
            "field_requests": df.requests,
            "closed": df.closed,
            "open_in_file": df.open_in_file,
            "p50_days": df.p50_hours / 24,
            "p90_days": df.p90_hours / 24,
            "benchmark_2019_p90_days": df.p90_hours_2019 / 24,
            "share_within_benchmark": df.share_within_2019_p90,
            "record_gap": df.record_gap.map({True: "gap", False: ""}),
        }
    )
    write_table(
        ws,
        df,
        formats={
            "field_requests": "#,##0",
            "closed": "#,##0",
            "open_in_file": "#,##0",
            "p50_days": "0.00",
            "p90_days": "0.00",
            "benchmark_2019_p90_days": "0.00",
            "share_within_benchmark": "0%",
        },
    )
    color_scale(ws, "K", 2, len(df) + 1)
    for i, gap in enumerate(df.record_gap, start=2):
        if gap:
            for j in range(1, df.shape[1] + 1):
                ws.cell(row=i, column=j).font = NOTE_FONT


def published_los(wb: Workbook) -> None:
    ws = wb.create_sheet("Published LOS")
    ws["A1"] = (
        "DOTI Level of Service goals (business days), denvergov.org, retrieved 2026-09-29. "
        "Goals are current; earlier years are a retrospective comparison. Open cases count as misses."
    )
    ws["A1"].font = NOTE_FONT
    df = export("published_los")
    df["record_gap"] = df.record_gap.map({True: "gap", False: ""})
    write_table(
        ws,
        df,
        start_row=2,
        formats={
            "requests": "#,##0",
            "closed": "#,##0",
            "share_within_los": "0%",
            "p50_business_days": "0",
            "p90_business_days": "0",
        },
    )
    color_scale(ws, "F", 3, len(df) + 2)


def by_agency(wb: Workbook) -> None:
    ws = wb.create_sheet("KPI by Agency")
    df = export("agency_year")
    df = pd.DataFrame(
        {
            "agency": df.agency,
            "year": df.year,
            "field_requests": df.field_requests,
            "p50_days": df.p50_hours / 24,
            "p90_days": df.p90_hours / 24,
            "share_open_in_file": df.share_open_in_file,
        }
    ).sort_values(["agency", "year"])
    write_table(
        ws,
        df,
        formats={
            "field_requests": "#,##0",
            "p50_days": "0.00",
            "p90_days": "0.00",
            "share_open_in_file": "0%",
        },
    )
    ws.conditional_formatting.add(
        f"C2:C{len(df) + 1}", DataBarRule(start_type="min", end_type="max", color="0F766E")
    )


def monthly(wb: Workbook) -> None:
    ws = wb.create_sheet("Monthly")
    df = export("monthly_category")
    wide = df.pivot_table(
        index="month", columns="request_category", values="field_requests", aggfunc="sum"
    ).fillna(0)
    wide = wide[wide.sum().sort_values(ascending=False).index].astype(int).reset_index()
    wide["month"] = pd.to_datetime(wide["month"]).dt.date
    n_cat = wide.shape[1] - 1
    write_table(ws, wide, formats={c: "#,##0" for c in wide.columns[1:]} | {"month": "yyyy-mm"})
    total_col = n_cat + 2
    roll_col = n_cat + 3
    for j, label in [(total_col, "total"), (roll_col, "rolling_12m")]:
        cell = ws.cell(row=1, column=j, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    first, last = get_column_letter(2), get_column_letter(n_cat + 1)
    total_letter = get_column_letter(total_col)
    for i in range(2, len(wide) + 2):
        ws.cell(
            row=i, column=total_col, value=f"=SUM({first}{i}:{last}{i})"
        ).number_format = "#,##0"
        if i >= 13:
            ws.cell(
                row=i, column=roll_col, value=f"=SUM({total_letter}{i - 11}:{total_letter}{i})"
            ).number_format = "#,##0"
    ws.column_dimensions[get_column_letter(total_col)].width = 10
    ws.column_dimensions[get_column_letter(roll_col)].width = 13
    ws.freeze_panes = "B2"


def channel(wb: Workbook) -> None:
    ws = wb.create_sheet("Channel")
    df = export("channel_year").sort_values(["year", "field_requests"], ascending=[True, False])
    write_table(
        ws,
        df,
        formats={
            "contacts": "#,##0",
            "field_requests": "#,##0",
            "share_of_contacts": "0.0%",
            "share_of_field": "0.0%",
        },
    )
    start = len(df) + 4
    ws.cell(
        row=start - 1, column=1, value="Median hours to close, phone vs app (closed field requests)"
    ).font = Font(bold=True, size=12)
    write_table(
        ws,
        export("channel_speed"),
        start_row=start,
        formats={
            "phone_closed": "#,##0",
            "app_closed": "#,##0",
            "phone_p50_hours": "0.0",
            "app_p50_hours": "0.0",
            "app_minus_phone_hours": "0.0",
        },
    )
    ws.freeze_panes = "A2"


def neighborhoods(wb: Workbook) -> None:
    ws = wb.create_sheet("Neighborhoods")
    df = (
        export("neighborhood")
        .drop(columns=["centroid_lat", "centroid_lon"])
        .sort_values("wait_index", ascending=False)
    )
    write_table(
        ws,
        df,
        formats={
            "population": "#,##0",
            "per_capita_income": "$#,##0",
            "pct_poverty": "0.0",
            "pct_renters": "0.0",
            "field_requests": "#,##0",
            "field_requests_2025": "#,##0",
            "field_requests_per_1k_per_year": "0.0",
            "share_open_in_file": "0%",
            "wait_index": "0.000",
            "ranked_requests": "#,##0",
            "p50_days": "0.00",
        },
    )
    col = get_column_letter(list(df.columns).index("wait_index") + 1)
    ws.conditional_formatting.add(
        f"{col}2:{col}{len(df) + 1}",
        ColorScaleRule(
            start_type="num",
            start_value=0.45,
            start_color="0F766E",
            mid_type="num",
            mid_value=0.5,
            mid_color="F8FAFC",
            end_type="num",
            end_value=0.55,
            end_color="1E3A5F",
        ),
    )


def main() -> None:
    wb = Workbook()
    for build in (
        readme,
        kpis,
        dictionary,
        quality,
        by_type,
        published_los,
        by_agency,
        monthly,
        channel,
        neighborhoods,
    ):
        build(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(
        f"wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size / 1024:.0f} KB, {len(wb.sheetnames)} tabs)"
    )


if __name__ == "__main__":
    main()
